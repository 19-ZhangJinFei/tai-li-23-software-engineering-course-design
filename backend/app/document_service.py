from __future__ import annotations
from io import BytesIO
from pathlib import Path
from docx import Document as WordDocument
from pypdf import PdfReader
try:
    from llama_index.core.node_parser import SentenceSplitter
except Exception:  # pragma: no cover - lightweight environments
    SentenceSplitter = None
from sqlalchemy import delete
from sqlalchemy.orm import Session
from .ai import get_ai
from .models import Document, DocumentChunk, now
from .storage import storage


ALLOWED_EXTENSIONS = {".docx", ".pdf", ".txt", ".md"}


def extract_pages(filename: str, data: bytes) -> list[tuple[int | None, str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".docx":
        doc = WordDocument(BytesIO(data))
        parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if text:
                    parts.append(text)
        return [(None, "\n".join(parts))]
    if suffix == ".pdf":
        reader = PdfReader(BytesIO(data))
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    if suffix in {".txt", ".md"}:
        return [(None, data.decode("utf-8", errors="replace"))]
    raise ValueError("不支持的文档格式")


def split_text(text: str) -> list[str]:
    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if not text:
        return []
    if SentenceSplitter:
        splitter = SentenceSplitter(chunk_size=600, chunk_overlap=100, separator="\n")
        return [chunk.strip() for chunk in splitter.split_text(text) if chunk.strip()]
    chunks, start = [], 0
    while start < len(text):
        chunks.append(text[start:start + 600])
        start += 500
    return chunks


def process_document(db: Session, document: Document) -> Document:
    document.status = "processing"
    document.error_message = None
    db.commit()
    try:
        pages = extract_pages(document.original_name, storage.read(document.storage_key))
        prepared: list[tuple[int | None, str]] = []
        for page_number, text in pages:
            prepared.extend((page_number, chunk) for chunk in split_text(text))
        if not prepared:
            raise ValueError("文档未提取到有效文本")
        ai = get_ai()
        embeddings: list[list[float]] = []
        for i in range(0, len(prepared), 10):
            embeddings.extend(ai.embed([text for _, text in prepared[i:i + 10]]))
        db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == document.id))
        for idx, ((page_number, content), embedding) in enumerate(zip(prepared, embeddings)):
            db.add(DocumentChunk(
                document_id=document.id,
                course_id=document.course_id,
                chunk_index=idx,
                content=content,
                page_number=page_number,
                metadata_json={"filename": document.original_name},
                embedding=embedding,
            ))
        document.status = "ready"
        document.chunk_count = len(prepared)
        document.processed_at = now()
        db.commit()
        db.refresh(document)
        return document
    except Exception as exc:
        db.rollback()
        document.status = "failed"
        document.error_message = str(exc)[:1000]
        db.add(document)
        db.commit()
        raise

