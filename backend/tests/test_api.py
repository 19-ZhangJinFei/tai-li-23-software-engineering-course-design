from io import BytesIO
from docx import Document as WordDocument


def make_docx() -> bytes:
    doc = WordDocument()
    doc.add_heading("RAG 基础", level=1)
    doc.add_paragraph("RAG 是检索增强生成技术，主要包含索引、检索和生成三个阶段。")
    doc.add_paragraph("Embedding 将文本转换为表示语义的数值向量，向量数据库用于相似内容检索。")
    doc.add_paragraph("知识库没有可靠依据时，系统应当拒绝回答，避免大语言模型产生幻觉。")
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def get_demo_course(client, headers):
    response = client.get("/api/v1/courses", headers=headers)
    assert response.status_code == 200
    return response.json()["data"]["items"][0]


def test_health_and_auth(client):
    assert client.get("/health").json()["data"]["status"] == "healthy"
    response = client.post("/api/v1/auth/register", json={
        "email": "new.student@example.com", "display_name": "新同学", "password": "Password@123"
    })
    assert response.status_code == 201
    assert response.json()["data"]["status"] == "active"
    duplicate = client.post("/api/v1/auth/register", json={
        "email": "new.student@example.com", "display_name": "新同学", "password": "Password@123"
    })
    assert duplicate.status_code == 409
    bad_login = client.post("/api/v1/auth/login", json={
        "email": "new.student@example.com", "password": "wrong"
    })
    assert bad_login.status_code == 401


def test_course_permissions(client, student_headers, admin_headers):
    forbidden = client.post("/api/v1/courses", headers=student_headers,
                            json={"code": "NEW", "title": "无权创建"})
    assert forbidden.status_code == 403
    created = client.post("/api/v1/courses", headers=admin_headers,
                          json={"code": "TEST-COURSE", "title": "测试课程", "description": "测试"})
    assert created.status_code == 201
    joined = client.post("/api/v1/courses/join", headers=student_headers, json={"code": "TEST-COURSE"})
    assert joined.status_code == 200


def test_document_rag_summary_quiz_flow(client, admin_headers, student_headers):
    course = get_demo_course(client, admin_headers)
    payload = make_docx()
    upload = client.post(
        "/api/v1/documents/upload-local",
        headers=admin_headers,
        data={"course_id": course["id"]},
        files={"file": ("rag-test.docx", payload, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload.status_code == 201, upload.text
    document = upload.json()["data"]
    process = client.post(f"/api/v1/documents/{document['id']}/process", headers=admin_headers)
    assert process.status_code == 200, process.text
    assert process.json()["data"]["status"] == "ready"
    assert process.json()["data"]["chunk_count"] > 0

    conv = client.post("/api/v1/conversations", headers=student_headers,
                       json={"course_id": course["id"]}).json()["data"]
    answer = client.post(f"/api/v1/conversations/{conv['id']}/messages", headers=student_headers,
                         json={"content": "RAG 包含哪些阶段？"})
    assert answer.status_code == 200
    assert "event: done" in answer.text
    messages = client.get(f"/api/v1/conversations/{conv['id']}/messages", headers=student_headers).json()["data"]
    assert len(messages) == 2
    assert messages[-1]["citations"]

    for unrelated in ("红烧肉应该放几勺糖？", "明天太原的天气怎么样？", "莫扎特出生在哪一年？"):
        refusal = client.post(f"/api/v1/conversations/{conv['id']}/messages", headers=student_headers,
                              json={"content": unrelated})
        assert "当前课程资料中未找到可靠依据" in refusal.text

    summary = client.post(f"/api/v1/courses/{course['id']}/summary", headers=student_headers)
    assert summary.status_code == 200
    assert summary.json()["data"]["key_points"]

    quiz = client.post(f"/api/v1/courses/{course['id']}/quizzes", headers=student_headers,
                       json={"count": 5})
    assert quiz.status_code == 201
    quiz_id = quiz.json()["data"]["id"]
    quiz_detail = client.get(f"/api/v1/quizzes/{quiz_id}", headers=student_headers).json()["data"]
    assert len(quiz_detail["questions"]) == 5
    assert "correct_index" not in quiz_detail["questions"][0]
    result = client.post(f"/api/v1/quizzes/{quiz_id}/submit", headers=student_headers,
                         json={"answers": [0, 0, 0, 0, 0]})
    assert result.status_code == 200
    assert result.json()["data"]["score"] == 100.0


def test_admin_dashboard(client, admin_headers):
    response = client.get("/api/v1/dashboard/admin", headers=admin_headers)
    assert response.status_code == 200
    assert response.json()["data"]["users"] >= 2


def test_course_data_isolation(client, admin_headers, student_headers):
    private = client.post("/api/v1/courses", headers=admin_headers,
                          json={"code": "PRIVATE-AI", "title": "未加入课程"}).json()["data"]
    denied = client.get(f"/api/v1/courses/{private['id']}/documents", headers=student_headers)
    assert denied.status_code == 403
    assert denied.json()["code"] == "FORBIDDEN"
